import openpyxl
import pandas as pd
import io

# Test 1: openpyxl loading Excel file with 14-digit string/integer
wb = openpyxl.Workbook()
ws = wb.active
ws['A1'] = "Enrollment Number"
ws['A2'] = 24002170210063 # integer in excel
ws['A3'] = "24002170210063" # string in excel

buf = io.BytesIO()
wb.save(buf)
buf.seek(0)

# Load with pandas read_excel
df1 = pd.read_excel(buf, dtype=str)
print("pd.read_excel dtype=str:")
print(df1['Enrollment Number'].tolist())

buf.seek(0)
# Load with openpyxl data_only=True
wb2 = openpyxl.load_workbook(buf, data_only=True)
ws2 = wb2.active
print("\nopenpyxl cell A2 value type:", type(ws2['A2'].value), repr(ws2['A2'].value))
print("openpyxl cell A3 value type:", type(ws2['A3'].value), repr(ws2['A3'].value))
